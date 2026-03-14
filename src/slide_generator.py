"""
PPTX slide generator with HBS Africa Business Club branding.

Slide order:
  1. Title slide  (dark navy + first background image)
  2. N × Takeaway slides  (uploaded background images, cycling)
  3. 3 × Quote slides  (dark navy)

All slides are 10″ × 10″ (square) for Instagram / LinkedIn.
"""

from pathlib import Path

from lxml import etree  # type: ignore
from PIL import Image  # type: ignore
from pptx import Presentation  # type: ignore
from pptx.dml.color import RGBColor  # type: ignore
from pptx.enum.text import PP_ALIGN  # type: ignore
from pptx.oxml.ns import qn  # type: ignore
from pptx.util import Inches, Pt  # type: ignore

# ---------------------------------------------------------------------------
# Brand palette
# ---------------------------------------------------------------------------
NAVY  = RGBColor(0x1B, 0x2A, 0x4A)
GREEN = RGBColor(0x9D, 0xC5, 0x28)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LGRAY = RGBColor(0xCC, 0xCC, 0xCC)

SLIDE_W = Inches(10)
SLIDE_H = Inches(10)
MARGIN  = Inches(0.65)
CONTENT_W = SLIDE_W - 2 * MARGIN
BOTTOM_BAR_H = Inches(0.52)


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _blank_layout(prs):
    return prs.slide_layouts[6]


def _send_to_back(slide, shape):
    sp_tree = slide.shapes._spTree
    sp_tree.remove(shape._element)
    sp_tree.insert(2, shape._element)


def _solid_bg(slide, color: RGBColor):
    rect = slide.shapes.add_shape(1, 0, 0, SLIDE_W, SLIDE_H)
    rect.fill.solid()
    rect.fill.fore_color.rgb = color
    rect.line.fill.background()
    _send_to_back(slide, rect)


def _set_alpha(shape, opacity_pct: int):
    spPr = shape._element.spPr
    solid = spPr.find(".//" + qn("a:solidFill"))
    if solid is None:
        return
    clr = solid.find(qn("a:srgbClr")) or solid.find(qn("a:schemeClr"))
    if clr is None:
        return
    for a in clr.findall(qn("a:alpha")):
        clr.remove(a)
    alpha_el = etree.SubElement(clr, qn("a:alpha"))
    alpha_el.set("val", str(int(opacity_pct * 1000)))


def _add_bg_image(slide, image_path: str):
    pic = slide.shapes.add_picture(str(image_path), 0, 0, SLIDE_W, SLIDE_H)
    _send_to_back(slide, pic)


def _add_overlay(slide, opacity_pct: int = 65):
    rect = slide.shapes.add_shape(1, 0, 0, SLIDE_W, SLIDE_H)
    rect.fill.solid()
    rect.fill.fore_color.rgb = NAVY
    rect.line.fill.background()
    _set_alpha(rect, opacity_pct)


def _add_bottom_bar(slide, label: str = "HBS Africa Business Club"):
    bar = slide.shapes.add_shape(1, 0, SLIDE_H - BOTTOM_BAR_H, SLIDE_W, BOTTOM_BAR_H)
    bar.fill.solid()
    bar.fill.fore_color.rgb = GREEN
    bar.line.fill.background()

    tb = slide.shapes.add_textbox(
        MARGIN, SLIDE_H - BOTTOM_BAR_H + Inches(0.1),
        SLIDE_W - 2 * MARGIN, Inches(0.35),
    )
    tf = tb.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = label
    run.font.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = NAVY


def _add_top_stripe(slide):
    bar = slide.shapes.add_shape(1, 0, 0, SLIDE_W, Inches(0.12))
    bar.fill.solid()
    bar.fill.fore_color.rgb = GREEN
    bar.line.fill.background()


def _textbox(slide, text, left, top, width, height,
             size, color, bold=False, align=PP_ALIGN.LEFT, wrap=True):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = wrap
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    return tb


# ---------------------------------------------------------------------------
# Image pre-processing
# ---------------------------------------------------------------------------

def preprocess_image(src_path: str, out_path: str, size: int = 2160) -> str:
    with Image.open(src_path) as img:
        if img.mode == "RGBA":
            bg = Image.new("RGB", img.size, (27, 42, 74))
            bg.paste(img, mask=img.split()[3])
            img = bg
        elif img.mode != "RGB":
            img = img.convert("RGB")

        w, h = img.size
        s = min(w, h)
        img = img.crop(((w - s) // 2, (h - s) // 2, (w + s) // 2, (h + s) // 2))
        img = img.resize((size, size), Image.LANCZOS)
        img.save(out_path, "JPEG", quality=88)
    return out_path


# ---------------------------------------------------------------------------
# Slide builders
# ---------------------------------------------------------------------------

def _title_slide(prs, bg_path, panel_name, event_name, event_date):
    slide = prs.slides.add_slide(_blank_layout(prs))

    if bg_path:
        _add_bg_image(slide, bg_path)
        _add_overlay(slide, opacity_pct=70)
    else:
        _solid_bg(slide, NAVY)

    _add_top_stripe(slide)

    _textbox(slide, "HBS AFRICA BUSINESS CLUB",
             MARGIN, Inches(0.25), CONTENT_W, Inches(0.4),
             size=12, color=GREEN, bold=True)

    _textbox(slide, "Key Takeaways",
             MARGIN, Inches(1.1), CONTENT_W, Inches(0.7),
             size=28, color=GREEN, bold=True)

    _textbox(slide, panel_name or "Panel Discussion",
             MARGIN, Inches(1.85), CONTENT_W, Inches(3.2),
             size=36, color=WHITE, bold=True)

    bottom_text = event_name + (f"  ·  {event_date}" if event_date else "")
    _textbox(slide, bottom_text,
             MARGIN, Inches(6.0), CONTENT_W, Inches(0.5),
             size=14, color=LGRAY)

    _add_bottom_bar(slide, event_name)


def _takeaway_slide(prs, takeaway: dict, bg_path: str, event_name: str):
    slide = prs.slides.add_slide(_blank_layout(prs))

    if bg_path:
        _add_bg_image(slide, bg_path)
        _add_overlay(slide, opacity_pct=68)
    else:
        _solid_bg(slide, NAVY)

    _add_top_stripe(slide)

    _textbox(slide, "KEY TAKEAWAY",
             MARGIN, Inches(0.22), CONTENT_W, Inches(0.35),
             size=10, color=GREEN, bold=True)

    num_tb = slide.shapes.add_textbox(MARGIN, Inches(0.65), Inches(0.55), Inches(0.55))
    p = num_tb.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = str(takeaway.get("number", ""))
    run.font.size = Pt(22)
    run.font.color.rgb = GREEN
    run.font.bold = True

    _textbox(slide, takeaway.get("headline", ""),
             MARGIN + Inches(0.55), Inches(0.58), CONTENT_W - Inches(0.55), Inches(1.4),
             size=21, color=WHITE, bold=True)

    div = slide.shapes.add_shape(1, MARGIN, Inches(2.1), CONTENT_W, Inches(0.04))
    div.fill.solid()
    div.fill.fore_color.rgb = GREEN
    div.line.fill.background()

    y = Inches(2.25)
    sub_bullets = takeaway.get("sub_bullets", [])
    available_h = SLIDE_H - BOTTOM_BAR_H - Inches(0.2) - y
    per_bullet_h = min(Inches(1.5), available_h / max(len(sub_bullets), 1))

    for bullet in sub_bullets:
        if y + per_bullet_h > SLIDE_H - BOTTOM_BAR_H - Inches(0.15):
            break
        tb = slide.shapes.add_textbox(
            MARGIN + Inches(0.1), y,
            CONTENT_W - Inches(0.1), per_bullet_h,
        )
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.space_before = Pt(0)
        p.space_after = Pt(0)

        bold_text = bullet.get("bold_phrase", "")
        if bold_text:
            r1 = p.add_run()
            r1.text = bold_text + ".  "
            r1.font.size = Pt(13)
            r1.font.bold = True
            r1.font.color.rgb = GREEN

        body_text = bullet.get("body", "")
        if body_text:
            r2 = p.add_run()
            r2.text = body_text
            r2.font.size = Pt(12)
            r2.font.bold = False
            r2.font.color.rgb = WHITE

        y += per_bullet_h

    _add_bottom_bar(slide, event_name)


def _quote_slide(prs, quote: dict, event_name: str):
    slide = prs.slides.add_slide(_blank_layout(prs))
    _solid_bg(slide, NAVY)

    top_bar_h = Inches(0.55)
    bar = slide.shapes.add_shape(1, 0, 0, SLIDE_W, top_bar_h)
    bar.fill.solid()
    bar.fill.fore_color.rgb = GREEN
    bar.line.fill.background()

    _textbox(slide, "MEMORABLE QUOTE",
             MARGIN, Inches(0.14), CONTENT_W, Inches(0.3),
             size=11, color=NAVY, bold=True)

    accent = slide.shapes.add_shape(
        1, Inches(0.35), top_bar_h + Inches(0.3),
        Inches(0.1), SLIDE_H - top_bar_h - BOTTOM_BAR_H - Inches(0.6),
    )
    accent.fill.solid()
    accent.fill.fore_color.rgb = GREEN
    accent.line.fill.background()

    _textbox(slide, "\u201C",
             MARGIN, top_bar_h + Inches(0.1), Inches(1.2), Inches(1.2),
             size=90, color=GREEN, bold=True)

    _textbox(slide, quote.get("text", ""),
             MARGIN + Inches(0.15), top_bar_h + Inches(1.05),
             CONTENT_W - Inches(0.15), Inches(5.8),
             size=22, color=WHITE)

    speaker = quote.get("speaker", "")
    if speaker and speaker.lower() not in ("unknown", ""):
        attr_tb = slide.shapes.add_textbox(
            MARGIN + Inches(0.15), Inches(7.9),
            CONTENT_W - Inches(0.15), Inches(0.6),
        )
        p = attr_tb.text_frame.paragraphs[0]
        r1 = p.add_run()
        r1.text = "\u2014 "
        r1.font.size = Pt(15)
        r1.font.bold = True
        r1.font.color.rgb = GREEN
        r2 = p.add_run()
        r2.text = speaker
        r2.font.size = Pt(15)
        r2.font.bold = True
        r2.font.color.rgb = WHITE

    _add_bottom_bar(slide, event_name)


# ---------------------------------------------------------------------------
# Public entry point
# ---------------------------------------------------------------------------

def generate_pptx(
    analysis: dict,
    bg_images: list,
    output_path: str,
    panel_name: str = "",
    event_name: str = "HBS Africa Business Conference",
    event_date: str = "",
    work_dir: str = "/tmp",
) -> str:
    prs = Presentation()
    prs.slide_width  = SLIDE_W
    prs.slide_height = SLIDE_H

    processed = []
    for i, src in enumerate(bg_images):
        dst = str(Path(work_dir) / f"bg_{i}.jpg")
        processed.append(preprocess_image(src, dst))

    def _bg(index: int):
        return processed[index % len(processed)] if processed else None

    _title_slide(prs, _bg(0), panel_name, event_name, event_date)

    for i, takeaway in enumerate(analysis.get("takeaways", [])):
        _takeaway_slide(prs, takeaway, _bg(i), event_name)

    for quote in analysis.get("quotes", []):
        _quote_slide(prs, quote, event_name)

    prs.save(output_path)
    return output_path
